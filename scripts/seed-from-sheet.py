#!/usr/bin/env python3
"""
Builds the question-bank seed from Nihad's question workbook.

    PYTHONIOENCODING=utf-8 python scripts/seed-from-sheet.py

writes two files:

    supabase/seeds/all-questions.sql   paste into the Supabase SQL editor and run
    supabase/seeds/REVIEW.md           what Nihad still has to check or supply

Rerun it whenever the workbook changes (new options, a fixed answer, a passage
added) and paste the new SQL: every question is matched on its paper and
number, so a rerun updates rows instead of adding copies.

Where things come from
----------------------
* The workbook (`--sheet`): one row per question, both papers in one sheet,
  told apart by the Date column. Header on row 1:
      Q#, Subject, Topic, Subtopic, Question (English), Difficulty,
      Answer Options, Correct Answer, Image, Date, Status, Source
  `Answer Options` holds one option per line ("A) ...\\nB) ..."). Lines that
  are not lettered options (numbered statements, "Words: ...") belong to the
  question and are added to its text.
* The older Russian maths workbook (`--ru-maths`): only its
  `Question (Russian)` column, the Russian wording of 12 April Q57-Q81.
* scripts/seed-from-sheet.toml: everything written by hand rather than read
  from a cell -- the Russian Language questions put into Russian, the Russian
  wording of the 26 April maths, marking guides, fixes, and the reason a
  question is held back as a draft. Edit that file, not the SQL.

Rows with a picture (an embedded image, or anything in the Image column) are
left out entirely: their diagrams are added through the admin form, and this
seed does not touch those rows in the database at all.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import tomllib
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA_FILE = Path(__file__).with_name("seed-from-sheet.toml")
OUT_SQL = ROOT / "supabase" / "seeds" / "all-questions.sql"
OUT_REVIEW = ROOT / "supabase" / "seeds" / "REVIEW.md"

DOWNLOADS = Path.home() / "Downloads"
DEFAULT_SHEET = DOWNLOADS / "Questions_12_and_26_April_2026_ENGLISH_ALL_OPTIONS.xlsm"
DEFAULT_RU_MATHS = DOWNLOADS / "Questions (RU maths).xlsm"

# The sheet's Date column -> the paper. `ref` is the external_ref prefix the
# 30 Sep maths seed already used for 12 April, so those rows are updated.
PAPERS = {
    "12 April 2026": {
        "ref": "dim-2026-04-12-variant-a",
        "source": "DİM Grade 9, Russian sector, Variant A, 12 April 2026",
        "label": "12 Apr (A)",
    },
    "26 April 2026": {
        "ref": "dim-2026-04-26-variant-d",
        "source": "DİM Grade 9, Russian sector, Variant D, 26 April 2026",
        "label": "26 Apr (D)",
    },
}
PAPER_YEAR = 2026

SUBJECTS = {"English": "english", "Russian Language": "russian", "Mathematics": "math"}

# Sheet topic -> the topic id the test builder lists. The starter ids in
# packages/shared (grammar, spelling, reading, writing, probability, ...) are
# reused wherever the sheet's topic means the same thing, so the builder does
# not show "Grammar: 0 questions" beside "Morphology"; the rest follow the
# 30 Sep maths seed (the sheet's name, lower-case, dashed). The sheet's own
# finer label is kept in `subtopic`.
TOPICS = {
    "english": {
        "Reading": "reading",
        "Grammar": "grammar",
        "Vocabulary": "vocabulary",
        "Writing": "writing",
    },
    "russian": {
        "Phonetics": "phonetics",
        "Morphology": "grammar",
        "Syntax": "grammar",
        "Orthography": "spelling",
        "Vocabulary": "vocabulary",
        "Phraseology": "vocabulary",
        "Reading": "reading",
        "Text Analysis": "reading",
        "Reading / Writing": "writing",
    },
    "math": {
        "Arithmetic": "arithmetic",
        "Algebra": "algebra",
        "Geometry": "geometry",
        "Coordinate Geometry": "coordinate-geometry",
        "Number Theory": "number-theory",
        "Sets & Logic": "sets-logic",
        "Probability": "probability",
        # The starter topic is "Probability and Statistics".
        "Statistics": "probability",
    },
}

DIFFICULTY = {1: "easy", 2: "medium", 3: "hard"}

LETTERS = "ABCDE"
LETTERED = re.compile(r"^([A-E])\)\s*(.*)$")
# "B", "B — 16%", "C - sentence 4": a letter, then optionally what it says.
LETTER_ANSWER = re.compile(r"^([A-E])(?:\s*[—–-]\s*(.*))?$")
# "1,4" or "1 and 5 — embarrassed; disrespectful": more than one statement.
MULTI_ANSWER = re.compile(r"^\d+\s*(?:,|and)\s*\d+")


@dataclass
class Question:
    paper: str
    number: int
    row: int
    subject: str  # subject_id
    topic: str
    subtopic: str
    difficulty: str
    type: str
    prompt: str
    options: list[str]
    answer: str
    letter: str | None
    status: str = "ready"
    translations: dict = field(default_factory=dict)
    missing: list[str] = field(default_factory=list)  # why it is not ready
    check: list[str] = field(default_factory=list)  # for Nihad to look over
    machine: list[str] = field(default_factory=list)  # what was translated by machine

    @property
    def ref(self) -> str:
        return f"{PAPERS[self.paper]['ref']}:q{self.number}"

    @property
    def source(self) -> str:
        return f"{PAPERS[self.paper]['source']}, Q{self.number}"


warnings: list[str] = []


def warn(message: str) -> None:
    warnings.append(message)


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return unicodedata.normalize("NFC", str(value)).strip()


def loose(value: str) -> str:
    """For comparing an answer's gloss with an option: no quotes, case, end dots."""
    value = re.sub(r"[«»“”\"'‘’]", "", value).casefold()
    return re.sub(r"\s+", " ", value).strip(" .;")


def header_map(ws) -> tuple[int, dict[str, int]]:
    for row in ws.iter_rows(min_row=1, max_row=10):
        if text(row[0].value) == "Q#":
            return row[0].row, {text(c.value): c.column - 1 for c in row if c.value}
    raise SystemExit(f"No header row starting with 'Q#' in sheet {ws.title!r}.")


def image_rows(ws) -> set[int]:
    return {img.anchor._from.row + 1 for img in getattr(ws, "_images", [])}


def split_options(raw: str) -> tuple[list[str], list[str]]:
    """Lettered options A-E, and every other line (part of the question)."""
    lettered: dict[str, str] = {}
    extra: list[str] = []
    for line in (line.strip() for line in raw.splitlines()):
        if not line or line.startswith("No answer choices"):
            continue
        match = LETTERED.match(line)
        if match and match.group(1) not in lettered:
            lettered[match.group(1)] = match.group(2).strip()
        else:
            extra.append(line)
    letters = "".join(lettered)
    if letters and letters != LETTERS[: len(letters)]:
        raise ValueError(f"options are lettered {letters}, not A, B, C...")
    return [lettered[letter] for letter in letters], extra


def strip_inline_options(prompt: str) -> str:
    """'... Choose: A) x; B) y ...' -> '... Choose:' -- the options live in options."""
    match = re.search(r"(?:^|(?<=\s))A\)\s", prompt)
    if match and re.search(r"(?:^|(?<=\s))B\)\s", prompt[match.start():]):
        return prompt[: match.start()].rstrip()
    return prompt


def with_extra_lines(prompt: str, extra: list[str]) -> str:
    """Statements the sheet put in the options column belong in the question."""
    flat = re.sub(r"\s+", " ", prompt)
    missing = [line for line in extra if re.sub(r"\s+", " ", line) not in flat]
    return prompt + ("\n" + "\n".join(missing) if missing else "")


def russian_form(value: str) -> str:
    """How a maths option or answer is written in Russian: 4,5 / 26 см / (4; −1)."""
    value = re.sub(r"(?<=\d)\.(?=\d)", ",", value)
    value = re.sub(r"(?<=\d) cm", " см", value)
    value = re.sub(r"(?<=\d) s$", " с", value)
    value = re.sub(r"\((−?[\d,]+), (−?[\d,]+)\)", r"(\1; \2)", value)
    return value


def load_ru_maths(path: Path) -> dict[int, str]:
    if not path.exists():
        warn(f"Russian maths workbook not found ({path}); 12 April maths has no Russian text.")
        return {}
    wb = openpyxl.load_workbook(path, read_only=False)
    ws = wb["Questions"]
    header_row, cols = header_map(ws)
    found = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        number, russian = text(row[cols["Q#"]]), text(row[cols["Question (Russian)"]])
        if number.isdigit() and russian:
            found[int(number)] = russian
    return found


def build(sheet: Path, ru_maths_path: Path, data: dict) -> tuple[list[Question], list[tuple[str, int, int, str]]]:
    wb = openpyxl.load_workbook(sheet)
    ws = wb["Questions"]
    header_row, cols = header_map(ws)
    pictures = image_rows(ws)
    ru_maths = load_ru_maths(ru_maths_path)

    def cell(row, name):
        return text(row[cols[name]])

    questions: list[Question] = []
    skipped: list[tuple[str, int, int, str]] = []
    used_overrides: set[str] = set()

    for row_number, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        if not cell(row, "Q#"):
            continue
        paper = cell(row, "Date")
        if paper not in PAPERS:
            warn(f"row {row_number}: unknown date {paper!r}, skipped")
            continue
        number = int(cell(row, "Q#"))

        if row_number in pictures or cell(row, "Image"):
            skipped.append((paper, number, row_number, cell(row, "Subject")))
            continue

        subject = SUBJECTS[cell(row, "Subject")]
        sheet_topic = cell(row, "Topic")
        if sheet_topic not in TOPICS[subject]:
            raise SystemExit(f"row {row_number}: no topic id for {subject} topic {sheet_topic!r}; add it to TOPICS.")
        difficulty = DIFFICULTY[int(float(cell(row, "Difficulty")))]

        options, extra = split_options(cell(row, "Answer Options"))
        prompt = with_extra_lines(strip_inline_options(cell(row, "Question (English)")), extra if options else [])
        raw_answer = cell(row, "Correct Answer")
        no_choices = cell(row, "Answer Options").lower()

        q = Question(
            paper=paper,
            number=number,
            row=row_number,
            subject=subject,
            topic=TOPICS[subject][sheet_topic],
            subtopic=cell(row, "Subtopic"),
            difficulty=difficulty,
            type="multiple-choice",
            prompt=prompt,
            options=options,
            answer="",
            letter=None,
        )

        letter_match = LETTER_ANSWER.match(raw_answer)
        if "open-ended" in no_choices or raw_answer.lower().startswith("open-ended"):
            q.type = "open-ended"
        elif options and letter_match:
            q.letter = letter_match.group(1)
        elif MULTI_ANSWER.match(raw_answer):
            # Numbered statements with no A-E list: pick more than one.
            q.options, q.answer = [], raw_answer
            q.status = "draft"
            q.missing.append(
                f"More than one right answer ({raw_answer}); the app marks one choice only. "
                "Rewrite as A-E options that each list statement numbers, like the other questions."
            )
        elif not options:
            q.type = "short-answer"
            q.answer = raw_answer
        else:
            raise SystemExit(f"row {row_number}: cannot read answer {raw_answer!r} against options {options}")

        # ---- hand-written content from the data file ----
        override = data.get(q.ref, {})
        if override:
            used_overrides.add(q.ref)
        for key in ("topic", "subtopic", "type"):
            if key in override:
                setattr(q, key, override[key])
        if "prompt" in override:
            q.prompt = override["prompt"].strip()
        if "options" in override:
            if q.options and len(override["options"]) != len(q.options):
                raise SystemExit(f"{q.ref}: {len(override['options'])} options in the data file, {len(q.options)} in the sheet")
            q.options = list(override["options"])
        if override.get("translated"):
            q.machine.append(override.get("translated_what", "question text and English options put into Russian"))
        if "answer" in override:
            q.answer = override["answer"]

        if q.letter:
            q.answer = q.options[LETTERS.index(q.letter)]
            gloss = letter_match.group(2) if letter_match else None
            if gloss and loose(gloss) != loose(q.answer) and not override.get("options"):
                warn(f"{q.ref}: answer {raw_answer!r} reads differently from option {q.letter} {q.answer!r}")
        if q.type == "open-ended":
            q.options = []
            q.answer = override.get("guide", "").strip()

        # Maths: base text English, Russian in translations.ru.
        if subject == "math":
            ru_prompt = override.get("ru_prompt") or (ru_maths.get(number) if paper == "12 April 2026" else None)
            if ru_prompt:
                ru: dict = {"prompt": ru_prompt.strip()}
                ru_options = override.get("ru_options") or [russian_form(o) for o in q.options]
                if ru_options != q.options:
                    ru["options"] = ru_options
                if q.type == "short-answer":
                    ru_answer = override.get("ru_answer") or russian_form(q.answer)
                    if ru_answer != q.answer:
                        ru["correctAnswer"] = ru_answer
                q.translations = {"ru": ru}
                if paper == "26 April 2026" and override.get("ru_prompt"):
                    q.machine.append("Russian version (shown on the Russian site) translated by machine")
            else:
                warn(f"{q.ref}: no Russian text for this maths question")

        # Russian Language questions are asked in Russian.
        if subject == "russian" and "prompt" not in override:
            q.status = "draft"
            q.missing.append("Not put into Russian yet.")
        if subject == "russian":
            english = [o for o in [q.prompt, *q.options] if re.search(r"[A-Za-z]{3,}", o)]
            if english:
                warn(f"{q.ref}: Russian Language text still has English in it: {english[0][:60]!r}")

        if override.get("status"):
            q.status = override["status"]
        if override.get("missing"):
            q.missing.append(override["missing"].strip())
        if override.get("check"):
            q.check.append(override["check"].strip())

        validate(q)
        questions.append(q)

    for ref in sorted(set(data) - used_overrides):
        warn(f"data file entry {ref} matches no loaded row (an image row, or a typo)")

    return questions, skipped


def validate(q: Question) -> None:
    if q.status not in ("draft", "ready"):
        raise SystemExit(f"{q.ref}: status {q.status!r}")
    if q.status == "draft" and not q.missing:
        raise SystemExit(f"{q.ref}: a draft needs a `missing` note saying why")
    if q.status != "ready":
        return
    problem = None
    if not q.prompt:
        problem = "no question text"
    elif not q.answer.strip():
        problem = "no answer (or marking guide)"
    elif q.type == "multiple-choice" and (len(q.options) < 2 or q.answer not in q.options):
        problem = "options do not include the answer"
    elif len(set(q.options)) != len(q.options):
        problem = "two options are the same"
    if problem:
        raise SystemExit(f"{q.ref} is marked ready but has {problem}")


# --------------------------------------------------------------------------
# SQL
# --------------------------------------------------------------------------


def sql_text(value: str | None) -> str:
    if value is None:
        return "null"
    if "\n" in value or "\\" in value:
        return "E'" + value.replace("\\", "\\\\").replace("'", "''").replace("\n", "\\n") + "'"
    return "'" + value.replace("'", "''") + "'"


def sql_json(value) -> str:
    # A standard (non-E) literal: JSON's own backslash escapes pass through.
    return "'" + json.dumps(value, ensure_ascii=False).replace("'", "''") + "'::jsonb"


def migration_0009() -> str:
    body = (ROOT / "supabase" / "migrations" / "0009_written_answers.sql").read_text(encoding="utf-8")
    return body.strip()


def counts(questions: list[Question]) -> dict:
    table: dict = {}
    for q in questions:
        table.setdefault((q.paper, q.subject), {}).setdefault(q.status, 0)
        table[(q.paper, q.subject)][q.status] += 1
    return table


SUBJECT_NAMES = {"english": "English", "russian": "Russian Language", "math": "Maths"}


def write_sql(questions: list[Question], skipped, sheet: Path) -> str:
    total = len(questions)
    ready = sum(q.status == "ready" for q in questions)
    table = counts(questions)
    lines_counts = []
    for paper in PAPERS:
        for subject in ("english", "russian", "math"):
            c = table.get((paper, subject), {})
            if c:
                lines_counts.append(
                    f"--     {PAPERS[paper]['label']:<11} {SUBJECT_NAMES[subject]:<17}"
                    f"{c.get('ready', 0):>3} ready, {c.get('draft', 0):>2} draft"
                )
    skipped_text = "; ".join(
        f"{PAPERS[p]['label']} Q" + ", Q".join(str(n) for _, n, _, _ in sorted(s for s in skipped if s[0] == p))
        for p in PAPERS
    )

    out: list[str] = []
    w = out.append
    w("-- ============================================================================")
    w("-- Exampeak: the DİM Grade 9 question bank, Russian sector. Two papers:")
    w("--   12 April 2026, Variant A and 26 April 2026, Variant D (Q5-Q81 of each).")
    w("-- ============================================================================")
    w("--")
    w("-- Generated by scripts/seed-from-sheet.py from Nihad's workbook")
    w(f"-- ({sheet.name}). Do not edit this file by hand:")
    w("-- change the workbook or scripts/seed-from-sheet.toml and rerun the script.")
    w("--")
    w("-- Paste the whole file into the Supabase SQL editor and run it once. It is one")
    w("-- transaction: if anything fails, nothing is changed.")
    w("--")
    w("-- What it does")
    w("--   1. Adds written answers marked by the AI marker (migration 0009): the")
    w("--      `open-ended` question type and the marker's feedback column.")
    w("--   2. Adds `questions.translations` (migration 0010): a question's text in")
    w("--      the student's other site language. Maths is stored in English with")
    w("--      its Russian wording here, so the Russian site shows it in Russian.")
    w(f"--   3. Loads {total} questions, {ready} of them ready for students:")
    out.extend(lines_counts)
    w("--      A draft is in the bank and on the admin page, but never in a test.")
    w("--      supabase/seeds/REVIEW.md says what each draft is waiting for.")
    w("--")
    w("-- Left out on purpose, and not touched in the database at all:")
    w(f"--   * questions with a picture ({skipped_text}). Their diagrams go in")
    w("--     through the admin form. The 12 April ones stay image-pending, as the")
    w("--     30 Sep maths seed left them.")
    w("--   * Q1-Q4 of each paper, the listening questions: there is no audio.")
    w("--")
    w("-- This supersedes dim-2026-04-12-variant-a/maths-questions.sql. It rewrites")
    w("-- the 12 April maths that seed loaded (Q57-Q71, Q73-Q76, Q80): the question")
    w("-- text becomes English, with the Russian wording moved into `translations`,")
    w("-- and the options and answers are filled in. Don't run the old file again.")
    w("--")
    w("-- Safe to run again. Each question is matched on its paper and number")
    w("-- (`external_ref`), so a second run updates the same rows rather than adding")
    w("-- copies. A rerun does put every column below back to what the sheet says,")
    w("-- so fix a question in the sheet (or the data file), not only in the admin")
    w("-- dashboard. Marks and pictures are never written here and are left alone.")
    w("--")
    w("-- Requires migrations 0001-0008 (supabase/run-all.sql) to have been run.")
    w("")
    w("begin;")
    w("")
    w("do $$")
    w("begin")
    w("  if to_regclass('public.questions') is null")
    w("     or not exists (select 1 from information_schema.columns")
    w("                    where table_schema = 'public' and table_name = 'questions'")
    w("                      and column_name in ('status'))")
    w("     or not exists (select 1 from information_schema.columns")
    w("                    where table_schema = 'public' and table_name = 'questions'")
    w("                      and column_name in ('external_ref')) then")
    w("    raise exception 'The questions table, or its status and external_ref columns, are missing. Run supabase/run-all.sql first.';")
    w("  end if;")
    w("end $$;")
    w("")
    w("-- ---------------------------------------------------------------------------")
    w("-- 1. Migration 0009: written answers")
    w("-- ---------------------------------------------------------------------------")
    w("")
    w(migration_0009())
    w("")
    w("-- ---------------------------------------------------------------------------")
    w("-- 2. Migration 0010: translations")
    w("-- ---------------------------------------------------------------------------")
    w("-- A question's text in the other site language, shaped")
    w('--   { "ru": { "prompt": "...", "options": [...], "correctAnswer": "..." } }')
    w("-- `options` is only there when an option reads differently (\"8 см³\"), in the")
    w("-- same order as the question's own; `correctAnswer` only for a short answer")
    w("-- that does. The full migration 0010 adds the same column and constraint")
    w("-- under the same names, so running both is harmless.")
    w("")
    w("alter table questions")
    w("  add column if not exists translations jsonb not null default '{}'::jsonb;")
    w("")
    w("alter table questions")
    w("  drop constraint if exists questions_translations_is_object;")
    w("")
    w("alter table questions")
    w("  add constraint questions_translations_is_object")
    w("  check (jsonb_typeof(translations) = 'object');")
    w("")
    w("-- ---------------------------------------------------------------------------")
    w("-- 3. The questions")
    w("-- ---------------------------------------------------------------------------")
    w("")
    w("-- Rows from an earlier load of 12 April that predate external_ref: give them")
    w("-- one, so the load below updates them instead of adding duplicates.")
    w("update questions")
    w("set external_ref = 'dim-2026-04-12-variant-a:q' || substring(source from ', Q([0-9]+)$')")
    w("where external_ref is null")
    w("  and source like 'DİM Grade 9, Russian sector, Variant A, 12 April 2026, Q%';")
    w("")
    w("insert into questions as q")
    w("  (external_ref, source, subject_id, topic_id, subtopic, difficulty, type, prompt,")
    w("   options, correct_answer, paper_year, status, translations)")
    w("select v.external_ref, v.source, v.subject_id, v.topic_id, v.subtopic, v.difficulty, v.type, v.prompt,")
    w("       v.options, v.correct_answer, v.paper_year, v.status, v.translations")
    w("from (values")

    rows = []
    current = None
    for q in questions:
        header = (q.paper, q.subject)
        if header != current:
            current = header
            rows.append(f"  -- {PAPERS[q.paper]['label']}, {SUBJECT_NAMES[q.subject]}")
        values = ", ".join(
            [
                sql_text(q.ref),
                sql_text(q.source),
                sql_text(q.subject),
                sql_text(q.topic),
                sql_text(q.subtopic),
                sql_text(q.difficulty),
                sql_text(q.type),
                sql_text(q.prompt),
                sql_json(q.options),
                sql_text(q.answer),
                str(PAPER_YEAR),
                sql_text(q.status),
                sql_json(q.translations),
            ]
        )
        rows.append(f"  ({values}),")
    # The last row has no trailing comma.
    for index in range(len(rows) - 1, -1, -1):
        if rows[index].endswith("),"):
            rows[index] = rows[index][:-1]
            break
    out.extend(rows)
    w(") as v(external_ref, source, subject_id, topic_id, subtopic, difficulty, type, prompt,")
    w("        options, correct_answer, paper_year, status, translations)")
    w("on conflict (external_ref) do update")
    w("set source         = excluded.source,")
    w("    subject_id     = excluded.subject_id,")
    w("    topic_id       = excluded.topic_id,")
    w("    subtopic       = excluded.subtopic,")
    w("    difficulty     = excluded.difficulty,")
    w("    type           = excluded.type,")
    w("    prompt         = excluded.prompt,")
    w("    options        = excluded.options,")
    w("    correct_answer = excluded.correct_answer,")
    w("    paper_year     = excluded.paper_year,")
    w("    status         = excluded.status,")
    w("    translations   = excluded.translations;")
    w("")
    w("commit;")
    w("")
    w("-- ---------------------------------------------------------------------------")
    w(f"-- Check: {total} rows from this file ({ready} ready), plus the 12 April")
    w("-- picture questions still image-pending from the 30 Sep seed.")
    w("-- ---------------------------------------------------------------------------")
    w("")
    w("select split_part(external_ref, ':', 1) as paper,")
    w("       subject_id,")
    w("       status,")
    w("       count(*) as questions,")
    w("       string_agg(split_part(external_ref, ':', 2), ', ' order by length(external_ref), external_ref) as which")
    w("from questions")
    w("where external_ref like 'dim-2026-04-12-variant-a:%'")
    w("   or external_ref like 'dim-2026-04-26-variant-d:%'")
    w("group by 1, 2, 3")
    w("order by 1, 2, 3;")
    return "\n".join(out) + "\n"


# --------------------------------------------------------------------------
# REVIEW.md
# --------------------------------------------------------------------------


def md(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def write_review(questions: list[Question], skipped) -> str:
    table = counts(questions)
    out: list[str] = []
    w = out.append
    w("# Question bank: what to check")
    w("")
    w("For Nihad. Generated by `scripts/seed-from-sheet.py` from your workbook, next to")
    w("`all-questions.sql`. Rerun the script after you change the sheet and this list")
    w("updates itself.")
    w("")
    w("**ready** means students can get it in a test. **draft** means it is saved but")
    w("hidden from students until the thing in the \"What's missing\" column is fixed.")
    w("")
    w("| Paper | Subject | Ready | Draft |")
    w("| --- | --- | ---: | ---: |")
    for paper in PAPERS:
        for subject in ("english", "russian", "math"):
            c = table.get((paper, subject))
            if c:
                w(f"| {PAPERS[paper]['label']} | {SUBJECT_NAMES[subject]} | {c.get('ready', 0)} | {c.get('draft', 0)} |")
    w("")
    w("## What you can do")
    w("")
    w("1. **Reading passages.** Many reading questions say \"according to the text\", but")
    w("   the texts themselves are not in the sheet, so students could only guess. Type")
    w("   each passage into the sheet (or send the scans) and those questions can go live.")
    w("2. **Questions with more than one right answer** (\"choose the true statements\"")
    w("   with no A-E list). The app marks one choice per question. Write A-E options")
    w("   that list statement numbers (\"A) 1, 4\"), the way the paper does elsewhere.")
    w("3. **Check the Russian.** Every Russian Language question was written in English")
    w("   in the sheet, so it was put into Russian by machine. Anything in «quotes» was")
    w("   kept exactly as you typed it; the rest needs your eye. Rows marked *yes* in")
    w("   the last column.")
    w("4. **Pictures.** Questions with a diagram are not loaded yet (list at the end).")
    w("   Add each picture through the admin page.")
    w("")
    w("## Every question")
    w("")
    w("| Paper | Q | Subject | Type | Status | What's missing / to check | Machine-translated? |")
    w("| --- | ---: | --- | --- | --- | --- | --- |")
    for q in questions:
        notes = [*q.missing, *q.check]
        w(
            f"| {PAPERS[q.paper]['label']} | {q.number} | {SUBJECT_NAMES[q.subject]} | {q.type} | "
            f"{q.status} | {md(' '.join(notes))} | {md('yes: ' + '; '.join(q.machine)) if q.machine else ''} |"
        )
    w("")
    w("## Not loaded: questions with a picture")
    w("")
    for paper in PAPERS:
        nums = sorted(n for p, n, _, _ in skipped if p == paper)
        w(f"- **{PAPERS[paper]['label']}**: Q" + ", Q".join(str(n) for n in nums))
    w("")
    w("The 12 April maths ones (Q72, Q77, Q78, Q79, Q81) are already in the bank as")
    w("\"image coming soon\" from the September load. Q1-Q4 (listening) are not loaded")
    w("because there is no audio.")
    w("")
    w("## How topics are named in the app")
    w("")
    w("Your sheet's topics are grouped into the topics students pick in the test")
    w("builder. Your exact topic is still saved with each question.")
    w("")
    for subject, mapping in TOPICS.items():
        pairs = ", ".join(f"{k} → `{v}`" for k, v in mapping.items())
        w(f"- **{SUBJECT_NAMES[subject]}**: {pairs}")
    return "\n".join(out) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--sheet", type=Path, default=DEFAULT_SHEET)
    parser.add_argument("--ru-maths", type=Path, default=DEFAULT_RU_MATHS)
    parser.add_argument("--out", type=Path, default=OUT_SQL)
    parser.add_argument("--review", type=Path, default=OUT_REVIEW)
    args = parser.parse_args()

    data = tomllib.loads(DATA_FILE.read_text(encoding="utf-8"))
    questions, skipped = build(args.sheet, args.ru_maths, data)

    args.out.write_text(write_sql(questions, skipped, args.sheet), encoding="utf-8", newline="\n")
    args.review.write_text(write_review(questions, skipped), encoding="utf-8", newline="\n")

    for message in warnings:
        print("warning:", message, file=sys.stderr)
    table = counts(questions)
    for (paper, subject), c in sorted(table.items()):
        print(f"{PAPERS[paper]['label']:<11} {SUBJECT_NAMES[subject]:<17} {c}")
    print(f"{len(questions)} questions, {len(skipped)} picture rows left out")
    print(f"wrote {args.out.relative_to(ROOT)} ({args.out.stat().st_size:,} bytes) and {args.review.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
